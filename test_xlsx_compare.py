"""Regression tests for comparison accuracy and safe report publication."""

import os
import sys
from contextlib import ExitStack
from functools import partial
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pandas as pd
import pytest
from mininterface import Mininterface, run
from openpyxl import Workbook, load_workbook
from pandas.io.excel._openpyxl import OpenpyxlWriter

import xlsx_compare as app


@pytest.fixture(autouse=True)
def noninteractive_cli(monkeypatch):
    """Parse real CLI arguments without opening a GUI or using saved configuration."""
    monkeypatch.setattr(
        app,
        "run",
        partial(run, interface=Mininterface, config_file=False, ask_for_missing=False),
    )


def make_workbook(path, sheets):
    """Create small inputs without relying on private workbooks."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, cells in sheets.items():
        sheet = workbook.create_sheet(name)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    workbook.save(path)
    workbook.close()
    return path


def run_cli(monkeypatch, left, right, output, *options):
    """Run the real CLI entrypoint with controlled arguments."""
    monkeypatch.setattr(sys, "argv", ["xlsx_compare.py", str(left), str(right), str(output), *options])
    return app.main()


def summary_rows(workbook):
    """Index summary rows by their original sheet names."""
    return {row[0]: row[1:] for row in workbook["COMPARISON"].iter_rows(min_row=5, values_only=True)}


@pytest.mark.parametrize("literal", ["NA", "NULL", "N/A"])
def test_literal_missing_value_strings(tmp_path, monkeypatch, literal):
    """Distinguish literal missing-value strings from empty cells."""
    left = make_workbook(tmp_path / "left.xlsx", {"Sheet": {"A1": literal}})
    right = make_workbook(tmp_path / "right.xlsx", {"Sheet": {}})
    output = tmp_path / "report.xlsx"
    assert run_cli(monkeypatch, left, right, output) == 0
    workbook = load_workbook(output)
    try:
        assert workbook["df-Sheet"]["A1"].value == f"{literal} -> "
        assert summary_rows(workbook)["Sheet"] == ("Differences found", "df-Sheet")
    finally:
        workbook.close()


def test_formulas_dimensions_coordinates_and_unique_sheets(tmp_path, monkeypatch):
    """Preserve formulas, coordinates, unique sheets, and summary formatting."""
    left = make_workbook(
        tmp_path / "left.xlsx",
        {
            "Sheet": {"A1": "=1+1"},
            "OnlyLeft": {},
            "Empty": {},
        },
    )
    right = make_workbook(
        tmp_path / "right.xlsx",
        {
            "Sheet": {"A1": "=2+2", "C3": "new"},
            "OnlyRight": {},
            "Empty": {},
        },
    )
    output = tmp_path / "report.xlsx"
    assert run_cli(monkeypatch, left, right, output) == 0
    workbook = load_workbook(output)
    try:
        assert workbook["df-Sheet"]["A1"].value == "=1+1 -> =2+2"
        assert workbook["df-Sheet"]["A1"].data_type == "s"
        assert workbook["df-Sheet"]["C3"].value == " -> new"
        assert workbook["df-Sheet"]["B2"].value is None
        summary = summary_rows(workbook)
        assert summary["Empty"][0] == "No differences"
        assert summary["OnlyLeft"][0] == "Only in file1"
        assert summary["OnlyRight"][0] == "Only in file2"
        summary_sheet = workbook["COMPARISON"]
        assert [cell.value for cell in summary_sheet[1]] == [
            "Sheet Name",
            "Status",
            "Difference Sheet",
        ]
        for cell in summary_sheet[1]:
            assert cell.font.bold
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "top"
            for side in (
                cell.border.left,
                cell.border.right,
                cell.border.top,
                cell.border.bottom,
            ):
                assert side.style == "thin"
        for row in summary_sheet.iter_rows(min_row=2):
            assert all(not cell.has_style for cell in row)
        for row in workbook["df-Sheet"]:
            assert all(not cell.has_style for cell in row)
    finally:
        workbook.close()


def test_identical_values_and_whitespace(tmp_path, monkeypatch):
    """Ignore surrounding whitespace and omit unchanged difference sheets."""
    left = make_workbook(tmp_path / "left.xlsx", {"Sheet": {"A1": "  text\t", "B2": 42}})
    right = make_workbook(tmp_path / "right.xlsx", {"Sheet": {"A1": "text", "B2": 42}})
    output = tmp_path / "report.xlsx"
    assert run_cli(monkeypatch, left, right, output) == 0
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["COMPARISON"]
        assert summary_rows(workbook)["Sheet"][0] == "No differences"
    finally:
        workbook.close()


def test_report_names_and_all_formula_like_text(tmp_path, monkeypatch):
    """Disambiguate truncated names and keep formula-like report values as text."""
    names = ["A" * 28 + "One", "A" * 28 + "Two", "=1+1"]
    monkeypatch.chdir(tmp_path)
    left = make_workbook(Path("=left.xlsx"), {name: {"A1": "=old"} for name in names})
    right = make_workbook(Path("=right.xlsx"), {name: {"A1": "=new"} for name in names})
    output = Path("report.xlsx")
    assert run_cli(monkeypatch, left, right, output) == 0
    workbook = load_workbook(output)
    try:
        rows = summary_rows(workbook)
        reports = [rows[name][1] for name in names]
        assert len({name.casefold() for name in reports}) == 3
        assert all(len(name) <= 31 for name in reports)
        for report in reports:
            assert workbook[report]["A1"].value == "=old -> =new"
        assert workbook["COMPARISON"]["B2"].value == "=left.xlsx"
        assert workbook["COMPARISON"]["B3"].value == "=right.xlsx"
        for sheet in workbook:
            assert all(cell.data_type != "f" for row in sheet for cell in row)
    finally:
        workbook.close()


def test_report_names_are_case_insensitive(tmp_path):
    """Allocate distinct report names when source names differ only in case."""
    with pd.ExcelWriter(tmp_path / "report.xlsx", engine="openpyxl") as writer:
        first = app.save_differences_to_excel(writer, "Alpha", pd.DataFrame([["first"]]))
        second = app.save_differences_to_excel(writer, "alpha", pd.DataFrame([["second"]]))
        assert first.casefold() != second.casefold()
        assert writer.book[first]["A1"].value == "first"
        assert writer.book[second]["A1"].value == "second"


@pytest.mark.parametrize("target", ["left", "right", "alias", "hardlink"])
def test_output_cannot_replace_input(tmp_path, monkeypatch, target):
    """Reject input paths and aliases without modifying either workbook."""
    left = make_workbook(tmp_path / "left.xlsx", {"Sheet": {"A1": "left"}})
    right = make_workbook(tmp_path / "right.xlsx", {"Sheet": {"A1": "right"}})
    before = (left.read_bytes(), right.read_bytes())
    (tmp_path / "subdir").mkdir()
    output = {
        "left": left,
        "right": right,
        "alias": tmp_path / "subdir" / ".." / "left.xlsx",
    }.get(target)
    if target == "hardlink":
        output = tmp_path / "link.xlsx"
        try:
            os.link(left, output)
        except OSError as error:
            pytest.skip(f"Hard links unavailable: {error}")
    assert run_cli(monkeypatch, left, right, output) == 1
    assert (left.read_bytes(), right.read_bytes()) == before
    assert not list(tmp_path.glob(".xlsx-compare-*"))


@pytest.mark.parametrize(
    "selection",
    [
        ("",),
        ("   ",),
        ("Sheet", ""),
        ("Missing",),
        ("Sheet", "Missing"),
        ("Sheet,Missing",),
    ],
)
def test_invalid_sheet_selections_preserve_output(tmp_path, monkeypatch, selection):
    """Reject invalid selections without replacing an existing report."""
    left = make_workbook(tmp_path / "left.xlsx", {"Sheet": {}})
    right = make_workbook(tmp_path / "right.xlsx", {"Sheet": {}})
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"existing report")
    assert run_cli(monkeypatch, left, right, output, "--sheets", *selection) == 1
    assert output.read_bytes() == b"existing report"
    assert not list(tmp_path.glob(".xlsx-compare-*"))


def test_selection_order_and_deduplication(tmp_path, monkeypatch):
    """Compare each selected sheet once in the requested order."""
    left = make_workbook(tmp_path / "left.xlsx", {"A": {}, "B": {}})
    right = make_workbook(tmp_path / "right.xlsx", {"A": {}, "B": {}})
    output = tmp_path / "report.xlsx"
    assert run_cli(monkeypatch, left, right, output, "--sheets", " B ", "A", "B") == 0
    workbook = load_workbook(output)
    try:
        assert list(summary_rows(workbook)) == ["B", "A"]
    finally:
        workbook.close()


def track_resources(monkeypatch):
    """Track workbook closure and all input/output streams through public methods."""
    tracked = SimpleNamespace(opened=[], closed=[], streams=[])
    real_load = app.load_workbook
    real_writer = app.pd.ExcelWriter

    def tracked_load(*args, **kwargs):
        tracked.streams.append(args[0])
        workbook = real_load(*args, **kwargs)
        tracked.opened.append(workbook)
        real_close = workbook.close

        def tracked_close():
            real_close()
            tracked.closed.append(workbook)

        monkeypatch.setattr(workbook, "close", tracked_close)
        return workbook

    def tracked_writer(stream, **kwargs):
        tracked.streams.append(stream)
        return real_writer(stream, **kwargs)

    monkeypatch.setattr(app, "load_workbook", tracked_load)
    monkeypatch.setattr(app.pd, "ExcelWriter", tracked_writer)
    return tracked


def simulate_failure(monkeypatch, failure, left, right):
    """Inject an input, comparison, save, or publication failure."""

    def fail(*_args, **_kwargs):
        raise OSError("simulated failure")

    if failure == "missing":
        left.unlink()
    elif failure == "corrupt":
        right.write_bytes(b"not an Excel archive")
    elif failure == "xml":
        with ZipFile(right) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        contents["xl/worksheets/sheet1.xml"] = b"<worksheet"
        with ZipFile(right, "w") as archive:
            for name, content in contents.items():
                archive.writestr(name, content)
    elif failure == "save":
        monkeypatch.setattr(OpenpyxlWriter, "_save", fail)
    elif failure == "replace":
        monkeypatch.setattr(app.os, "replace", fail)
    elif failure == "compare":
        monkeypatch.setattr(app, "compare_dataframes_cell_by_cell", fail)


@pytest.mark.parametrize("failure", ["success", "missing", "corrupt", "xml", "save", "replace", "compare"])
def test_failures_close_resources_and_preserve_report(tmp_path, monkeypatch, capsys, failure):
    """Close resources on every exit and preserve existing reports on failure."""
    left = make_workbook(tmp_path / "left.xlsx", {"Sheet": {"A1": "left"}})
    right = make_workbook(tmp_path / "right.xlsx", {"Sheet": {"A1": "right"}})
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"existing report")
    tracked = track_resources(monkeypatch)
    simulate_failure(monkeypatch, failure, left, right)
    assert run_cli(monkeypatch, left, right, output) == (0 if failure == "success" else 1)
    error = capsys.readouterr().err
    assert ("Error:" in error) == (failure != "success")
    if failure == "missing":
        assert str(left.resolve()) in error
        assert "working directory" in error
    if failure == "success":
        assert output.read_bytes() != b"existing report"
    else:
        assert output.read_bytes() == b"existing report"
    assert tracked.closed == list(reversed(tracked.opened))
    assert all(stream.closed for stream in tracked.streams)
    assert not list(tmp_path.glob(".xlsx-compare-*"))


def expected_sheet_differences(inputs, name):
    """Compute expected changed cells independently of the application helpers."""
    values = [list(source[name].iter_rows(values_only=True)) for source in inputs]
    expected = {}
    for i in range(max(map(len, values))):
        pair = [data[i] if i < len(data) else () for data in values]
        for j in range(max(map(len, pair))):
            cells = [row[j] if j < len(row) else None for row in pair]
            texts = ["" if value is None else str(value).strip() for value in cells]
            if texts[0] != texts[1]:
                expected[(i + 1, j + 1)] = f"{texts[0]} -> {texts[1]}"
    return expected


def assert_report_sheet(workbook, inputs, name, summary):
    """Verify the summary status and changed cells for one example sheet."""
    status, report = summary
    if name not in inputs[0].sheetnames:
        assert status == "Only in file2"
    elif name not in inputs[1].sheetnames:
        assert status == "Only in file1"
    else:
        expected = expected_sheet_differences(inputs, name)
        assert status == ("Differences found" if expected else "No differences")
        if expected:
            actual = {
                (cell.row, cell.column): cell.value
                for row in workbook[report]
                for cell in row
                if cell.value is not None
            }
            assert actual == expected


def test_example_report(tmp_path, monkeypatch):
    """Match the bundled example report against independently computed differences."""
    root = Path(__file__).resolve().parent
    left, right = root / "example/a.xlsx", root / "example/b.xlsx"
    output = tmp_path / "example-report.xlsx"
    assert run_cli(monkeypatch, left, right, output) == 0
    workbook = load_workbook(output)
    inputs = [load_workbook(path, read_only=True, data_only=False) for path in (left, right)]
    try:
        rows = summary_rows(workbook)
        assert set(rows) == set(inputs[0].sheetnames) | set(inputs[1].sheetnames)
        for name, summary in rows.items():
            assert_report_sheet(workbook, inputs, name, summary)
    finally:
        workbook.close()
        for source in inputs:
            source.close()


@pytest.mark.parametrize(
    "arguments, expected_output, expected_sheets",
    [
        (["a.xlsx", "b.xlsx"], "comparison_output.xlsx", []),
        (["a.xlsx", "b.xlsx", "c1.xlsx"], "c1.xlsx", []),
        (
            ["a.xlsx", "b.xlsx", "c1.xlsx", "--sheets", "Sheet1", "Sheet2"],
            "c1.xlsx",
            ["Sheet1", "Sheet2"],
        ),
    ],
)
def test_positional_output(arguments, expected_output, expected_sheets):
    """Parse real CLI arguments without opening an interactive interface."""
    env = run(
        app.Env,
        interface=Mininterface,
        config_file=False,
        ask_for_missing=False,
        args=arguments,
    ).env

    assert env.file1 == Path("a.xlsx")
    assert env.file2 == Path("b.xlsx")
    assert env.output == Path(expected_output)
    assert env.sheets == expected_sheets


@pytest.mark.parametrize("output_name", [None, "c1.xlsx"])
def test_main_writes_requested_report(tmp_path, monkeypatch, output_name):
    """Write a real report using the default or an explicit output path."""
    for filename, value in (("a.xlsx", "old"), ("b.xlsx", "new")):
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        workbook.active["A1"] = value
        workbook.save(tmp_path / filename)
        workbook.close()

    monkeypatch.chdir(tmp_path)
    arguments = ["xlsx_compare.py", "a.xlsx", "b.xlsx"]
    if output_name is not None:
        arguments.append(output_name)
    monkeypatch.setattr(sys, "argv", arguments)
    monkeypatch.setattr(
        app,
        "run",
        partial(run, interface=Mininterface, config_file=False, ask_for_missing=False),
    )

    with ExitStack() as resources:

        def open_input(*args, **kwargs):
            workbook = load_workbook(*args, **kwargs)
            resources.callback(workbook.close)
            return workbook

        monkeypatch.setattr(app, "load_workbook", open_input)
        app.main()

    output_path = tmp_path / (output_name or "comparison_output.xlsx")
    assert output_path.is_file()
    report = load_workbook(output_path)
    try:
        assert report.sheetnames == ["df-Sheet1", "COMPARISON"]
        assert report["df-Sheet1"]["A1"].value == "old -> new"
        assert report["COMPARISON"]["A5"].value == "Sheet1"
        assert report["COMPARISON"]["B5"].value == "Differences found"
    finally:
        report.close()


@pytest.mark.parametrize("sheet_name", ["Sheet 2", "Sheet,1"])
def test_sheet_names_with_spaces_or_commas(tmp_path, monkeypatch, sheet_name):
    """Treat each CLI list item as one sheet name, including spaces and commas."""
    left = make_workbook(tmp_path / "left.xlsx", {sheet_name: {}, "Other": {}})
    right = make_workbook(tmp_path / "right.xlsx", {sheet_name: {}, "Other": {}})
    output = tmp_path / "report.xlsx"
    assert run_cli(monkeypatch, left, right, output, "--sheets", sheet_name) == 0
    workbook = load_workbook(output)
    try:
        assert list(summary_rows(workbook)) == [sheet_name]
    finally:
        workbook.close()
